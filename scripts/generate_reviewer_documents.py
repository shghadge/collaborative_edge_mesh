#!/usr/bin/env python3
import argparse
import csv
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
LIGHT_GRAY = "F2F2F2"
DARK = "333333"


def _shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shade = properties.find(qn("w:shd"))
    if shade is None:
        shade = OxmlElement("w:shd")
        properties.append(shade)
    shade.set(qn("w:fill"), fill)


def _set_cell(cell, text: object, *, bold: bool = False, white: bool = False) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(8)
    if white:
        run.font.color.rgb = RGBColor(255, 255, 255)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def _table(document: Document, headers: list[str], rows: list[list[object]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for index, header in enumerate(headers):
        _set_cell(table.rows[0].cells[index], header, bold=True, white=True)
        _shade(table.rows[0].cells[index], BLUE)
    for row_index, row in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(row):
            _set_cell(cells[index], value)
            if row_index % 2:
                _shade(cells[index], LIGHT_GRAY)
    document.add_paragraph()


def _configure(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(9.5)
    styles["Normal"].font.color.rgb = RGBColor.from_string(DARK)
    for name, size, color in (
        ("Title", 22, BLUE),
        ("Heading 1", 15, BLUE),
        ("Heading 2", 12, BLUE),
        ("Heading 3", 10.5, DARK),
    ):
        styles[name].font.name = "Arial"
        styles[name].font.size = Pt(size)
        styles[name].font.color.rgb = RGBColor.from_string(color)


def _heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def _paragraph(
    document: Document, text: str, *, bold_prefix: str | None = None
) -> None:
    paragraph = document.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix) :])
    else:
        paragraph.add_run(text)


def _bullet(document: Document, text: str) -> None:
    document.add_paragraph(text, style="List Bullet")


def _fmt(value: object, digits: int = 2) -> str:
    if value in {None, ""}:
        return "—"
    if isinstance(value, float):
        return f"{value:.{digits}f}"
    return str(value)


def _load_csv(path: Path) -> list[dict[str, str]]:
    if not path.exists():
        return []
    with path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def _select(groups: list[dict], experiment_group: str) -> list[dict]:
    return [row for row in groups if row["experiment_group"] == experiment_group]


def _core_rows(groups: list[dict]) -> list[list[object]]:
    rows = []
    selected = sorted(
        _select(groups, "core"),
        key=lambda row: (
            row["scenario"],
            int(row["node_count"]),
            row["sync_mode"],
        ),
    )
    for row in selected:
        rows.append(
            [
                "Full state" if row["sync_mode"] == "full_state" else "True delta",
                row["node_count"],
                "Clean" if row["scenario"] == "baseline" else "50/50 heal",
                row["runs"],
                f"{row['accepted_events_total']}/{row['attempted_events_total']}",
                f"{row['fully_observed_events_total']}/{row['accepted_events_total']}",
                f"{row['semantic_converged_runs']}/{row['runs']}",
                f"{row['converged_runs']}/{row['runs']}",
                _fmt(row["convergence_time_s_median_converged"]),
                _fmt(row["write_latency_p95_ms_median"]),
                _fmt(row["bytes_per_accepted_event_median"]),
            ]
        )
    return rows


def _fault_rows(groups: list[dict]) -> list[list[object]]:
    selected = _select(groups, "shape") + _select(groups, "degraded")
    rows = []
    for row in sorted(
        selected,
        key=lambda item: (
            item["experiment_group"],
            item["partition_split"],
            item["sync_mode"],
        ),
    ):
        condition = (
            row["partition_split"]
            if row["experiment_group"] == "shape"
            else "300 ms / 10% loss"
        )
        rows.append(
            [
                "Full state" if row["sync_mode"] == "full_state" else "True delta",
                condition,
                row["runs"],
                f"{row['semantic_converged_runs']}/{row['runs']}",
                f"{row['converged_runs']}/{row['runs']}",
                _fmt(row["convergence_time_s_median_converged"]),
                _fmt(row["bytes_per_accepted_event_median"]),
            ]
        )
    return rows


def _retention_rows(groups: list[dict]) -> list[list[object]]:
    rows = []
    for row in _select(groups, "retention"):
        rows.append(
            [
                row["event_id_retention_limit"],
                row["delta_log_limit"],
                row["runs"],
                row["delta_pulls_gap_total"],
                f"{row['semantic_converged_runs']}/{row['runs']}",
                f"{row['converged_runs']}/{row['runs']}",
                _fmt(row["full_state_repair_bytes_median"]),
            ]
        )
    return rows


def _resource_rows(groups: list[dict]) -> list[list[object]]:
    rows = []
    for row in sorted(
        _select(groups, "resource"),
        key=lambda item: (float(item["docker_cpus"]), item["docker_memory"], item["sync_mode"]),
    ):
        rows.append(
            [
                "Full state" if row["sync_mode"] == "full_state" else "True delta",
                row["docker_cpus"],
                row["docker_memory"],
                row["runs"],
                f"{row['semantic_converged_runs']}/{row['runs']}",
                f"{row['converged_runs']}/{row['runs']}",
                _fmt(row["write_latency_p95_ms_median"]),
                _fmt(row["cpu_peak_percent_median"]),
                _fmt(row["memory_peak_mib_median"]),
            ]
        )
    return rows


def _load_table_rows(groups: list[dict]) -> list[list[object]]:
    rows = []
    for row in sorted(
        _select(groups, "load"),
        key=lambda item: (
            int(item["node_count"]),
            float(item["target_offered_rate_events_s"]),
        ),
    ):
        rows.append(
            [
                row["node_count"],
                row["target_offered_rate_events_s"],
                row["runs"],
                _fmt(row["accepted_throughput_events_s_median"]),
                _fmt(row["write_latency_p95_ms_median"]),
                _fmt(row["write_success_rate_median"], 3),
                f"{row['converged_runs']}/{row['runs']}",
            ]
        )
    return rows


def _centralized_rows(groups: list[dict]) -> list[list[object]]:
    rows = []
    for row in sorted(
        _select(groups, "centralized"),
        key=lambda item: (item["scenario"], int(item["node_count"])),
    ):
        rows.append(
            [
                row["node_count"],
                "Available" if row["scenario"] == "baseline" else "Partitioned",
                row["runs"],
                f"{row['accepted_events_total']}/{row['attempted_events_total']}",
                row["rejected_events_total"],
                _fmt(row["write_latency_p95_ms_median"]),
                _fmt(row["failed_request_latency_p95_ms_median"]),
            ]
        )
    return rows


def _response_block(
    document: Document, number: str, comment: str, response: str
) -> None:
    _heading(document, f"Comment {number}", 3)
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Reviewer comment. ")
    run.bold = True
    paragraph.add_run(comment)
    paragraph = document.add_paragraph()
    run = paragraph.add_run("Response. ")
    run.bold = True
    paragraph.add_run(response)


def _add_completed_response(
    document: Document, evidence: dict, audit: dict, packet_rows: list[dict[str, str]]
) -> None:
    overall = evidence["overall"]
    _heading(document, "Part I — Response to Reviewers")
    _paragraph(
        document,
        "We thank the reviewers for distinguishing the architecture's contribution from the limits of the original single-host experiment. The revision narrows its claims, adds a centralized comparison, specifies faults exactly, adds event-level and semantic correctness checks, and reports repeated-run performance and resource measurements. It does not claim physical-edge, multi-host, cloud, or energy validation; those remain limitations rather than promised results.",
    )

    _heading(document, "Reviewer #3", 2)
    _response_block(
        document,
        "3.1",
        "Add comparisons, larger-scale or real-world deployment evidence, higher-load scalability, security, deployment challenges, and writing improvements.",
        f"We added a small centralized client-server baseline and an increasing-rate study, retained the submitted 5/10/15-node scale, and report {overall['run_count']} completed runs. We also added a deployment-limitations paragraph and corrected the manuscript's grammar and overbroad claims. We did not have access to independent edge devices or cloud VMs, so the revision explicitly identifies the testbed as single-host containerized evidence and does not present it as real-world deployment validation. Security remains outside the evaluated contribution and is stated as such.",
    )

    _heading(document, "Reviewer #4", 2)
    _response_block(
        document,
        "4.1",
        "The evaluation uses local Docker containers rather than real edge devices.",
        "Agreed. We changed the platform description and all conclusions to say 'single-host containerized testbed.' CPU and memory quotas emulate constrained resources but do not reproduce separate kernels, heterogeneous hardware, geographic paths, wireless links, energy use, or failure independence. No physical-edge or multi-host claim is made.",
    )
    _response_block(
        document,
        "4.2",
        "The comparisons are limited to internal CRDT transfer configurations.",
        "We added one deliberately small centralized client-server baseline using the same HTTP payload and request schedule. The baseline tests the paper's central availability tradeoff: when its only endpoint is blocked, disconnected writes are rejected, whereas partitioned CRDT replicas continue accepting local writes and reconcile after healing. This is not presented as a complete consensus-system benchmark.",
    )
    _response_block(
        document,
        "4.3",
        "Correct the 'EXPIREMENT' heading.",
        "Corrected to 'EXPERIMENT AND IMPLEMENTATION.' We also removed the phrases 'zero data loss' and 'only viable.' Coverage is now stated as no missing accepted tracking IDs at the final observation cutoff only for configurations where the histories were complete, alongside a separate semantic-state result.",
    )

    _heading(document, "Reviewer #7", 2)
    responses = [
        (
            "7.1",
            "Validate on physical or emulated multi-host infrastructure.",
            "This could not be completed with the available single workstation and no cloud allocation. We therefore made the limitation explicit and excluded Floci/container rehearsal from VM evidence. The revised evidence does not satisfy a physical or multi-host validation claim.",
        ),
        (
            "7.2",
            "Specify partition, delay, loss, bandwidth, reordering, and churn parameters.",
            "The primary partition is a bidirectional iptables drop between the selected groups. Writes begin only after the rules are confirmed. The configured partition hold is 15 s; the artifact records apply/remove timestamps and actual duration. The degraded-link profile applies 300 ms delay, 100 ms jitter, 10% correlated packet loss, and 2 Mbit/s bandwidth with Linux netem. The study includes clean, 50/50, 70/30, and 80/20 cases. Reordering and node churn were not injected and are explicitly excluded.",
        ),
        (
            "7.3",
            "Report per-event histories, duplicates, transient loss, overwritten values, and semantic correctness.",
            f"Every accepted write carries a deterministic tracking ID. Each node records its first-seen timestamp and local gossip round. The evaluator reports per-node missing-ID lists, histories missing at cutoff, duplicate submissions, time to first remote replica, time to last replica, and semantic state equality independently of Merkle-root equality. Across the completed dataset, {overall['accepted_events_total']} writes were accepted, {overall['fully_observed_events_total']} accepted tracking histories were present at every required replica at cutoff, {overall['missing_events_total']} accepted tracking histories were incomplete at cutoff, and {overall['duplicate_submissions_total']} duplicate submissions were recorded. Tracking coverage and semantic state are reported separately; neither is presented as a universal zero-loss proof.",
        ),
        (
            "7.4",
            "Specify clock synchronization, clock skew, and equal-timestamp LWW behavior.",
            "LWW versions are now ordered lexicographically by (timestamp, writer_id, event_id). Exact timestamp ties therefore resolve deterministically without arrival-order dependence. The property suite includes equal-timestamp, merge-order, and injected-clock-skew cases. The Docker runs share one host clock; cross-host clock synchronization was not evaluated.",
        ),
        (
            "7.5",
            "Specify all four CRDT merge semantics and correct the PN-counter definition.",
            "The revision defines G-Counter merge as the component-wise maximum; PN-Counter as two G-Counter maps P and N merged component-wise with value ΣP−ΣN; LWW-Register as the maximum deterministic version tuple; and OR-Set as union of tagged adds and removals, with active values equal to add tags not present in removals. Property tests exercise commutativity, associativity, idempotence, duplicate replay, and full NodeState convergence.",
        ),
        (
            "7.6",
            "Explain sliding-window correctness after long partitions.",
            "The bounded window applies only to the recent event-ID list used for local/debug visibility and state-growth control. It does not bound the separate first-seen evaluation histories, and it does not delete G-Counter components, PN-Counter components, LWW versions, OR-Set tags/tombstones, or other materialized CRDT state. A low-retention/low-delta-log adversarial case forces a detected delta gap and checks full-state anti-entropy fallback after healing.",
        ),
        (
            "7.7",
            "Define true-delta generation, causal metadata, acknowledgement, and repair.",
            "Each mutation increments a node-local sequence and stores a delta-log entry. Gossip sends a Merkle summary and current sequence over UDP. A receiver with a divergent root requests /state/deltas?since=<known-sequence> over HTTP. The responder returns retained deltas when the requested sequence remains available; otherwise it marks a gap and returns full state. The receiver joins the response into local CRDT state. A received delta that changes state is recorded under the receiver's next local sequence for onward relay; duplicate or cyclic delivery is not re-recorded because the idempotent merge makes no change. Thus dropped UDP summaries do not carry unique state: later summaries re-expose divergence and HTTP anti-entropy repairs it. The repeated clean and partition/heal matrix exercises this transitive repair path; timeout-qualified results are reported in the evidence tables rather than assumed from the mechanism.",
        ),
        (
            "7.8",
            "Report independent repetitions and statistical uncertainty.",
            f"The main matched full-state/true-delta matrix uses five deterministic seeds per configuration. The complete matrix contains {audit['expected_case_count']} declared case IDs and {audit['completed_case_count']} completed case IDs. The reporting artifacts include n, mean, median, standard deviation, quartiles, nonparametric bootstrap 95% intervals, exact paired sign-randomization p-values, and paired Cohen's dz. No p99 claim is made from five repetitions.",
        ),
        (
            "7.9",
            "Separate network payload, metadata, headers, repair, polling, retransmission, and serialization costs.",
            "Application counters now separate gossip summaries sent/received, full-state repair, delta repair, CRDT semantic state, evaluator metadata, centralized requests, and centralized responses. A representative packet capture separately reports Ethernet/IP/transport framing and repeated TCP payload segments. Serialization remains embedded in the corresponding application payload and is not claimed as an independently measured wire category.",
        ),
        (
            "7.10",
            "Measure throughput, propagation latency, resources, storage, and energy under increasing event rates.",
            "The added load matrix reports attempted and accepted rate, HTTP write latency, first-remote and last-replica propagation, recovery/convergence, CPU, peak CPU, RSS, network-interface bytes, block I/O, serialized semantic-state size, and evaluator-metadata size under 1-CPU/512-MiB limits, plus focused constrained-resource runs. Edge state is in memory, so serialized state size and block I/O are reported rather than implying a persistent edge-database footprint. Direct energy measurement was unavailable, so the revision removes energy-efficiency claims rather than estimating per-container energy from a shared host.",
        ),
    ]
    for number, comment, response in responses:
        _response_block(document, number, comment, response)

    if packet_rows:
        total_packets = sum(int(row["packet_count"]) for row in packet_rows)
        total_frames = sum(int(row["frame_bytes"]) for row in packet_rows)
        _paragraph(
            document,
            f"Packet-capture note. The representative captures contain {total_packets} classified protocol packets and {total_frames} captured frame bytes; the detailed protocol rows are retained with the raw capture artifacts.",
        )


def _add_paper_insert(
    document: Document,
    evidence: dict,
    audit: dict,
    packet_rows: list[dict[str, str]],
    figures_dir: Path | None,
) -> None:
    groups = evidence["groups"]
    overall = evidence["overall"]
    core_groups = _select(groups, "core")
    full_state_core = [row for row in core_groups if row["sync_mode"] == "full_state"]
    true_delta_core = [row for row in core_groups if row["sync_mode"] == "true_delta"]
    full_state_converged = sum(row["converged_runs"] for row in full_state_core)
    full_state_runs = sum(row["runs"] for row in full_state_core)
    full_state_semantic = sum(row["semantic_converged_runs"] for row in full_state_core)
    true_delta_converged = sum(row["converged_runs"] for row in true_delta_core)
    true_delta_runs = sum(row["runs"] for row in true_delta_core)
    true_delta_semantic = sum(row["semantic_converged_runs"] for row in true_delta_core)
    _heading(document, "Part II — Manuscript-Ready Evaluation Material")

    _heading(document, "Experimental scope and controls", 2)
    _paragraph(
        document,
        "The revised evaluation remains anchored to the submitted Python implementation: four state-based CRDTs, UDP gossip summaries, Merkle divergence detection, HTTP anti-entropy repair, fan-out three, a five-second gossip interval, and a ten-second gateway polling interval. Nodes run as isolated Docker containers on one x86-64 WSL2 host. Each node is limited to one CPU and 512 MiB unless a constrained-resource case states otherwise. The platform is an implementation-level systems testbed, not physical-edge or multi-host validation.",
    )
    _paragraph(
        document,
        "The principal matrix compares full-state and true-delta anti-entropy with 5, 10, and 15 nodes under clean operation and bidirectional 50/50 partition/heal. Each matched configuration uses seeds 7301–7305. Supplemental cases cover 70/30 and 80/20 partitions, a degraded 300 ms ± 100 ms/10% correlated-loss/2 Mbit/s link, increasing offered rates, a centralized reference, low delta retention, and CPU/memory limits. Runs are independent at the experiment level; events and packets are not treated as independent statistical samples.",
    )
    _paragraph(
        document,
        "A partition is active only after iptables rules are confirmed. The configured hold is 15 s, and every row preserves group membership plus fault-apply, fault-remove, and actual-duration timestamps. Healing removes only the network rules; no gateway merged-state push is used. Convergence is the first evaluator observation with one Merkle root, no missing accepted tracking IDs, and equal normalized semantic state. Because the evaluator polls every two seconds, convergence_time_s is a poll-quantized upper bound. Node first-seen timestamps provide separate propagation latency.",
    )

    _heading(document, "Primary repeated-run results", 2)
    _table(
        document,
        [
            "Method",
            "Nodes",
            "Condition",
            "n",
            "Accepted",
            "Seen by all",
            "Semantic",
            "Converged",
            "Median conv. s",
            "Median p95 write ms",
            "Median bytes/event",
        ],
        _core_rows(groups),
    )
    _paragraph(
        document,
        "Interpretation rule: 'seen by all' is accepted tracking-ID coverage at the final cutoff. 'Semantic' is normalized CRDT-state equality. 'Converged' additionally requires one Merkle root and both signals within the configured wait budget. A timeout is not recoded as eventual convergence, and incomplete tracking history is not silently relabeled as semantic data loss.",
    )
    _paragraph(
        document,
        f"Full-state repair satisfied the combined convergence criterion within the observation budget in {full_state_converged}/{full_state_runs} principal runs and semantic agreement in {full_state_semantic}/{full_state_runs}; corrected true-delta repair did so in {true_delta_converged}/{true_delta_runs} and {true_delta_semantic}/{true_delta_runs}, respectively. The mechanisms' median bytes per accepted event are reported separately. Transitive delta relay restores liveness but can forward state fragments through intermediate replicas, so the revision does not assume that true-delta repair is uniformly faster or smaller than full-state repair.",
    )
    _heading(document, "Partition shape and degraded-link results", 2)
    _table(
        document,
        [
            "Method",
            "Condition",
            "n",
            "Semantic",
            "Converged",
            "Median conv. s",
            "Median bytes/event",
        ],
        _fault_rows(groups),
    )

    if packet_rows:
        _heading(document, "Representative packet-layer accounting", 2)
        _table(
            document,
            [
                "Capture",
                "Protocol category",
                "Packets",
                "Captured frame bytes",
                "Transport payload bytes",
                "Framing/transport overhead",
                "Repeated TCP payload bytes",
            ],
            [
                [
                    Path(row["capture"]).name,
                    row["protocol"],
                    row["packet_count"],
                    row["frame_bytes"],
                    row["transport_payload_bytes"],
                    row["link_network_transport_overhead_bytes"],
                    row["repeated_tcp_payload_bytes"],
                ]
                for row in packet_rows
            ],
        )
        _paragraph(
            document,
            "The TCP category contains all node HTTP traffic on port 8000, including workload requests, evaluator/gateway observation, and anti-entropy repair. Endpoint-specific repair bytes come from the application counters; the packet layer is not relabeled as repair traffic. Repeated TCP payload is a sequence-number-based retransmission indicator within the capture window.",
        )

    _heading(document, "Increasing-rate behavior", 2)
    _table(
        document,
        [
            "Nodes",
            "Offered events/s",
            "n",
            "Median accepted/s",
            "Median p95 write ms",
            "Median success",
            "Converged",
        ],
        _load_table_rows(groups),
    )

    _heading(document, "Delta-gap fallback", 2)
    _table(
        document,
        [
            "Event-ID limit",
            "Delta-log limit",
            "n",
            "Detected gaps",
            "Semantic",
            "Converged",
            "Median full-repair bytes",
        ],
        _retention_rows(groups),
    )

    _heading(document, "Constrained-resource behavior", 2)
    _table(
        document,
        [
            "Method",
            "CPU limit",
            "Memory limit",
            "n",
            "Semantic",
            "Converged",
            "Median p95 write ms",
            "Median peak CPU %",
            "Median peak RSS MiB",
        ],
        _resource_rows(groups),
    )

    _heading(document, "Centralized availability reference", 2)
    _table(
        document,
        [
            "Clients",
            "Endpoint",
            "n",
            "Accepted/attempted",
            "Rejected",
            "Median accepted p95 ms",
            "Median failed p95 ms",
        ],
        _centralized_rows(groups),
    )
    _paragraph(
        document,
        "This reference isolates one tradeoff rather than ranking complete architectures. A single available endpoint avoids replica convergence; blocking that endpoint removes write availability. The CRDT mesh continues accepting writes on both sides of a partition and pays reconciliation cost after healing.",
    )

    _heading(document, "Correctness and statistics", 2)
    _paragraph(
        document,
        f"The artifact audit found {audit['completed_case_count']} completed case IDs out of {audit['expected_case_count']} enabled cases, {audit['result_row_count']} result rows, no duplicate case IDs, and no unexpected case IDs. Across all evaluation groups, {overall['accepted_events_total']} writes were accepted and {overall['fully_observed_events_total']} accepted tracking histories were present at all required replicas at cutoff; {overall['missing_events_total']} histories were incomplete at cutoff. Semantic-state equality is reported separately. These counts support bounded statements about the tested runs, not a proof of zero loss in all executions.",
    )
    _paragraph(
        document,
        "The CRDT test suite checks join commutativity, associativity, and idempotence; duplicate replay; PN-Counter positive/negative component merging; OR-Set add/remove semantics; deterministic LWW equal-timestamp and clock-skew behavior; full NodeState merge-order convergence; delta equivalence; and gap-triggered full-state fallback. The statistical report gives descriptive distributions and paired method differences. Tail claims are limited by the repetition count.",
    )

    if figures_dir:
        figure_specs = (
            ("availability_partition.png", "Write availability under partition"),
            ("convergence_by_scale.png", "Convergence success and observed time"),
            ("communication_cost.png", "Protocol bytes per accepted event"),
            ("load_response.png", "Increasing-rate throughput and latency"),
        )
        available = [
            (figures_dir / filename, caption)
            for filename, caption in figure_specs
            if (figures_dir / filename).exists()
        ]
        if available:
            _heading(document, "Revision figures", 2)
        for path, caption in available:
            document.add_picture(str(path), width=Inches(6.8))
            paragraph = document.add_paragraph(caption)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    _heading(document, "Limitations", 2)
    for limitation in (
        "All performance runs share one host, kernel, storage system, clock source, and Docker daemon. They do not establish geographically distributed, heterogeneous, wireless, physical-edge, or cloud behavior.",
        "CPU quotas and memory limits constrain containers but do not emulate ARM instruction behavior, independent power domains, separate kernels, or device-specific thermal throttling.",
        "No direct energy meter or host RAPL attribution was available; energy consumption is not reported.",
        "The degraded-link profile includes delay, jitter, correlated loss, and bandwidth limiting. It does not include packet reordering, one-way failure, or node churn.",
        "Application byte counters and representative packet captures provide category accounting, but JSON serialization bytes are part of their enclosing payload rather than a separate wire-layer category.",
        "The centralized reference tests endpoint availability only. It is not a benchmark against a replicated consensus datastore or a production cloud service.",
    ):
        _bullet(document, limitation)

    _heading(document, "Required manuscript claim corrections", 2)
    _bullet(
        document,
        "Abstract: replace 'zero data loss' with the measured accepted-event coverage and timeout-qualified convergence statement. Do not state that every method converged.",
    )
    _bullet(
        document,
        "Section III: describe gossip as a design choice for decentralized dissemination, not the 'only viable' mechanism. State that cloud export is architectural future work and was not implemented or evaluated.",
    )
    _bullet(
        document,
        "Section IV heading: replace 'EXPIREMENT AND IMPLEMENTATION' with 'EXPERIMENT AND IMPLEMENTATION.' Replace the PN-Counter and LWW merge rows with the complete deterministic definitions given in the reviewer response.",
    )
    _bullet(
        document,
        "Section V: replace the original single-aggregate Figure 3 and its four-method ranking with the repeated-run figures and primary table in this evidence package. Historical naïve/optimized-full-state numbers must not be mixed into the new matched full-state/true-delta analysis.",
    )
    _bullet(
        document,
        f"Section V finding: state that full-state repair converged within budget in {full_state_converged}/{full_state_runs} principal runs and corrected true-delta repair in {true_delta_converged}/{true_delta_runs}. Report the observed bytes per accepted event separately; do not claim in advance that either method is uniformly faster or smaller.",
    )
    _bullet(
        document,
        "Conclusion: limit the evidence to one single-host containerized implementation. Remove claims of physical-edge suitability, universal no-loss behavior, energy efficiency, cloud validation, and guaranteed true-delta superiority.",
    )

    _heading(document, "Artifact map", 2)
    for artifact in (
        "scripts/reviewer_experiment_matrix.csv — immutable case definitions and seeds",
        "results/reviewer_revision/**/eval-*.csv and eval-*.json — raw run rows and manifests",
        "results/reviewer_revision/summary/summary_statistics.csv — descriptive distributions",
        "results/reviewer_revision/summary/paired_comparisons.csv — paired differences, exact randomization tests, and effect sizes",
        "results/reviewer_revision/summary/run_audit.json — expected/completed/duplicate/unexpected case audit",
        "results/reviewer_revision/paper_evidence/paper_results_table.csv — paper-facing grouped values",
        "results/reviewer_revision/packet_capture/ — representative PCAP and layer-accounting rows",
        "results/reviewer_revision/figures/ — publication-oriented figures",
    ):
        _bullet(document, artifact)


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Generate completed reviewer response."
    )
    parser.add_argument("--evidence", type=Path, required=True)
    parser.add_argument("--audit", type=Path, required=True)
    parser.add_argument("--figures-dir", type=Path)
    parser.add_argument("--packet-summary", type=Path)
    parser.add_argument("--out", type=Path, required=True)
    args = parser.parse_args()

    evidence = json.loads(args.evidence.read_text(encoding="utf-8"))
    audit = json.loads(args.audit.read_text(encoding="utf-8"))
    packet_rows = _load_csv(args.packet_summary) if args.packet_summary else []

    document = Document()
    _configure(document)
    title = document.add_heading(
        "Response to Reviewers and Revised Evaluation Evidence", 0
    )
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = document.add_paragraph(
        "Paper 297 — Design and Evaluation of CRDT-driven Microservices at the Edge for Disaster Context Awareness"
    )
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph()
    _add_completed_response(document, evidence, audit, packet_rows)
    document.add_page_break()
    _add_paper_insert(document, evidence, audit, packet_rows, args.figures_dir)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    document.save(args.out)
    print(args.out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
